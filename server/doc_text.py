import os
import time
import json
import common
from common import SRC
import docx
from google.cloud import storage
from google.oauth2 import service_account
import py7zr
import config

bucket_name = 'liberman-transcriptions'
credentials = service_account.Credentials.from_service_account_info(json.loads(config.credential))
client = storage.Client(credentials=credentials, project=credentials.project_id)

source = 'DocText'
file_txt = 'liberman.txt'
t_begin = None


def load_file_bucket(blob_name):
    try:
        bucket = storage.Bucket(client, bucket_name)  # указать текущий начальный bucket
        temp_dir = common.current_path
        file_name = temp_dir + '/' + str(blob_name) + '.7z'
        if not os.path.exists(file_name):
            blob = bucket.blob(blob_name + '.7z')
            os.makedirs(temp_dir, exist_ok=True)
            blob.download_to_filename(file_name)  # это загрузка файла
        with py7zr.SevenZipFile(file_name, 'r') as arch:
            arch.extractall(path=temp_dir)
        os.remove(file_name)
        file_name = temp_dir + '/' + str(blob_name) + '.docx'
        return file_name
    except Exception as er:
        file_name = blob_name + '.7z' if blob_name is not None else 'None' + '" в облаке недоступен или отсутствует'
        txt = f'ERROR: {er}\n' + file_name
        common.write_log_db(
            '❌error', SRC, f"load_file_bucket\nОшибка {txt}", law_id=source)


def get_list_file():
    url = 'v2/select/{schema}/nsi_transcripts?column_order=date'.format(schema=config.schema_name)
    ans, is_ok, status = common.send_rest(url, params={"columns": "id, date"})
    if not is_ok:
        common.write_log_db(
            '❌error', SRC, f"get_list_file\nОшибка {ans}", law_id=source)
        return
    ans = json.loads(ans)
    return ans


def work():
    list_id = get_list_file()
    if list_id:
        for i, unit in enumerate(list_id):
            file_name = load_file_bucket('en_{date}'.format(date=unit['date']))
            result = ''  # текстовый документ
            if file_name:
                t1 = time.time()
                doc = docx.Document(file_name)
                for data in doc.paragraphs:
                    result += data.text + '\n'
                f = open(file_txt, 'a', encoding='utf-8')
                f.write(result)
                f.close()
                os.remove(file_name)
                print(i + 1, unit['date'], 'кол-во параграфов=' + common.str1000(len(doc.paragraphs)),
                      'td=' + common.get_duration(time.time() - t1),
                      'всего td=' + common.get_duration(time.time() - t_begin),
                      'size=' + common.str1000(os.path.getsize(file_txt)))


# common.current_path = os.path.abspath(os.curdir)
# t_begin = time.time()
# open(file_txt, 'w').close()  # если нужно создать и очистить
# work()
# print('Finish', common.get_duration(time.time() - t_begin))
